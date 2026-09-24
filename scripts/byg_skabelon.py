#!/usr/bin/env python3
"""Bygger vaerksdata_skabelon.xlsx (v4).

Ændringer fra v3 (gennemgang 24. september 2026, før mødet med Billund):
  * Afleveringsfristen er en dato: onsdag 7. oktober, ikke "uge 41".
  * DMI-område: ringsted er tilføjet til dropdown og hjælpetekst, så værker
    øst for Storebælt har et område. KRÆVER, at sysapp og df-data har
    området, og at config.KENDTE_DMI_OMRAADER og kontrollen i
    vaerksark_til_yaml.py følger med — ellers afviser konverteringen arket.
  * Tarif: arket siger nu, at båndenes tidspunkter er faste i modellen
    (N1's inddeling), og beder om tidspunkterne i mailen, hvis netselskabet
    bruger andre. Noterne pr. bånd beskriver de timer, konverteringen
    faktisk bruger.
  * Balancemarked: arket siger, at et 'ja' ved gasmotoren endnu ikke bruges.
    balancing._eligible_units_for_market lader kun elforbrugende enheder byde.
  * Eksempeltanken tank_lille er rettet fra 192 til 105 MWh, så begge
    eksempeltanke passer med hjælpetekstens 30 K.
  * NYT ark 'Varmepumpe' (Johns forslag): varmeproduktion og eloptag ved
    fuld last ved to-tre udetemperaturer pr. varmepumpe. COP ved 0 °C i
    Anlaeg!F er nu kun et nødtal, hvis arket er tomt, og maks varme må stå
    tom for varmepumper med målepunkter. Eksemplet er skaleret fra Billunds
    målte forhold (COP 2,4 / 2,9 / 3,4 ved -10 / 0 / +16 °C).

Ændringer fra v2 til v3:
Ændringer fra v2 (kollegernes tilbagemelding 17.-18. september 2026):
  * Timedata: årsproduktion ab værk i GWh/år som eget felt i B3. Feltet er
    altid påkrævet. Er der ingen timeserie, er det eneste, modellen har.
  * Anlaeg: kolonne G er elvirkningsgrad, ikke el-forhold (alpha). Alpha
    udledes i konverteringen — for varmepumper af COP, for elkedler af
    varmevirkningsgraden, for gasmotorer af el- og varmevirkningsgrad.
  * Anlaeg: tanke beskrives ved volumen, maks fyldning (MWh) og lade-/
    afladeeffekt. Delta T er ude af arket og udledes i konverteringen.
  * Priser: CO2 opgives i kr/t CO2, ikke kr/MWh gas. Modellen ganger selv
    med emissionsfaktoren 0,2 t CO2/MWh gas. Enheden i v2 var forkert.
  * Steens tekstrettelser: fast periode jul25-jun26 i stedet for "de seneste
    12 måneder", huller meldes i mailen og ikke på dagen.

Bevarede cellepositioner (vaerksark_til_yaml.py afhænger af dem):
  Timedata!B2 tidszone, overskrift i række 5, data fra række 6.
  Anlaeg: overskrift i række 4, enheder i række 5-23, 13 kolonner.
  Anlaeg: tankoverskrift i række 51, tanke i række 52-57, 5 kolonner.
  Priser: brændsler 6-10, faste led 14-17, bånd 22-24, B28/B29/B30.
NYT felt: Timedata!B3.
NYT ark (v4): Varmepumpe, overskrift i række 5, målepunkter fra række 6,
  kolonne A-D (navn, udetemperatur, varme, el). E er en COP-formel.
"""

from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ARIAL = "Arial"
H1 = Font(name=ARIAL, size=14, bold=True, color="1F3864")
H2 = Font(name=ARIAL, size=11, bold=True, color="1F3864")
BODY = Font(name=ARIAL, size=10)
BODY_B = Font(name=ARIAL, size=10, bold=True)
BODY_I = Font(name=ARIAL, size=10, italic=True, color="595959")
MONO = Font(name="Consolas", size=10, color="C00000")
HDR = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
INPUT_F = Font(name=ARIAL, size=10, color="0000FF")

HDR_FILL = PatternFill("solid", fgColor="1F3864")
UDFYLD = PatternFill("solid", fgColor="FFFF00")
EKSEMPEL = PatternFill("solid", fgColor="F2F2F2")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")

from openpyxl.styles import Side
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ENHEDSTYPER = ["heat_pump", "electric_boiler", "biomass_boiler", "gas_boiler",
               "gas_engine_chp", "solar_thermal", "waste_heat"]


def sat(ws, celle, vaerdi, font=BODY, fill=None, wrap=False, border=False):
    c = ws[celle]
    c.value = vaerdi
    c.font = font
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if border:
        c.border = BOX
    return c


def header_raekke(ws, raekke, kolonner, bredder):
    for i, (navn, bredde) in enumerate(zip(kolonner, bredder), start=1):
        c = ws.cell(row=raekke, column=i, value=navn)
        c.font = HDR
        c.fill = HDR_FILL
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = bredde
    ws.row_dimensions[raekke].height = 34


def dropdown(ws, vaerdier, omraade, titel):
    """Inline-liste. Samme form som DMI/priszone i v1, der virkede.

    v1's fejl var ikke formen, men dækningen: typelisten sad kun på de gule
    rækker, ikke på eksempelrækkerne, som er dem, man klikker på først.
    Derfor lægges den her på HELE enhedsblokken, eksempler inklusive."""
    dv = DataValidation(type="list", formula1='"' + ",".join(vaerdier) + '"',
                        allow_blank=True, showDropDown=False)
    dv.showErrorMessage = True
    dv.errorTitle = titel
    dv.error = "Vælg en værdi fra listen: " + ", ".join(vaerdier)
    dv.showInputMessage = True
    dv.promptTitle = titel
    dv.prompt = "Vælg fra listen"
    ws.add_data_validation(dv)
    dv.add(omraade)


wb = Workbook()

# ==============================================================================
# ARK 1 — VEJLEDNING
# ==============================================================================
ws = wb.active
ws.title = "Vejledning"
ws.sheet_view.showGridLines = False
for col, br in zip("ABCD", (3, 26, 64, 3)):
    ws.column_dimensions[col].width = br

sat(ws, "B2", "Dit værk i modellen — dataskabelon", H1)
sat(ws, "B3", "AI til driftsoptimering i fjernvarmen · 22. oktober 2026 · "
              "Fjernvarmens Hus, Kolding", BODY_I)

sat(ws, "B5", "Hvad du skal gøre inden kurset", H2)
opgaver = [
    ("1. Timedata",
     "Bed din SRO-leverandør eller driftsansvarlige om timeværdier for den "
     "samlede varmeproduktion ab værk, summeret over alle enheder, for "
     "1. juli 2025 til 30. juni 2026. Klokkeslæt i UTC, så sommertiden kan "
     "håndteres. Indsæt dem i arket 'Timedata'. Det er den ene ting med "
     "leveringstid — bestil den først."),
    ("2. Anlægsdata",
     "Udfyld arket 'Anlaeg' med dine produktionsenheder og "
     "akkumuleringstanke. Har I varmepumper, så udfyld også arket "
     "'Varmepumpe'."),
    ("3. Priser og tarif",
     "Udfyld arket 'Priser'. Tarifbåndene står på dit netselskabs prisblad. "
     "Brændselspriserne er dine egne indkøbspriser ekskl. moms."),
    ("4. Send arket til mig",
     "Send det udfyldte ark til mig senest onsdag 7. oktober, så jeg kan "
     "kontrollere det og bygge modelfilen inden kurset. Tag også filen med på "
     "dagen. Du skal ikke installere noget."),
]
for i, (hvad, tekst) in enumerate(opgaver, start=6):
    sat(ws, f"B{i}", hvad, BODY_B)
    sat(ws, f"C{i}", tekst, BODY, wrap=True)
    ws.row_dimensions[i].height = 58 if i == 6 else 44

sat(ws, "B11", "Tidszone — læs det her, også selvom resten springes over", H2)
sat(ws, "C12",
    "Modellen regner i UTC, og elpriser og vejrdata hentes i UTC. Derfor beder "
    "vi om dine timedata i UTC. Får du udtrækket i dansk tid, så konvertér IKKE "
    "selv — skriv i stedet 'dansk lokaltid' i tidszonefeltet øverst i arket "
    "Timedata, så klarer konverteringen det. En forkert tidszone flytter hele "
    "året en eller to timer, uden at noget ser forkert ud: varmen topper bare "
    "på det forkerte tidspunkt i forhold til elprisen, og modellen lærer et "
    "mønster, der ikke findes.", BODY, wrap=True)
ws.row_dimensions[12].height = 86

sat(ws, "B14", "Sådan får du timedata hjem", H2)
sat(ws, "C15",
    "Spørg driften eller din SRO-leverandør: »Jeg skal bruge et udtræk med "
    "timeværdier for den samlede varmeproduktion ab værk i MW for perioden "
    "1. juli 2025 til 30. juni 2026, som CSV eller Excel med to kolonner: "
    "tidsstempel og værdi. Tidsstempler i UTC, gerne som ISO 8601. Kan I kun "
    "levere lokal tid, så oplys det, og fortæl hvordan sommertidsskiftet ser "
    "ud.« Bed samtidig om den samlede årsproduktion ab værk i GWh — den skal "
    "stå i arket, uanset om timedataene kommer.",
    BODY, wrap=True)
ws.row_dimensions[15].height = 86

sat(ws, "B17", "Hvis varmeproduktionen har huller", H2)
sat(ws, "C18",
    "Enkelte manglende timer er normalt. Lad cellen stå tom — skriv ikke nul, "
    "for nul er en gyldig måling og bliver læst som 'værket producerede intet'. "
    "Mangler mere end en uge i træk, så skriv det i mailen, når du sender "
    "arket — ikke først på dagen. Der findes en syntesevej, hvor modellen "
    "danner de manglende timer ud fra DMI-vejrdata, men den skal sættes op "
    "inden kurset. Kommer der slet ingen timeserie, kan vi køre på "
    "årsproduktionen alene; så er resultatet et regneeksempel og ikke jeres "
    "drift.",
    BODY, wrap=True)
ws.row_dimensions[18].height = 86

sat(ws, "B20", "Varmepumper — arket 'Varmepumpe'", H2)
sat(ws, "C21",
    "En luft/vand-varmepumpe leverer mindre varme, når det er koldt, og mere, "
    "når det er lunt. Derfor beder vi ikke om ét tal, men om varmeproduktion "
    "og eloptag ved fuld last ved to-tre udetemperaturer, fx en kold dag "
    "(omkring −10 °C), omkring 0 °C og en varm dag (omkring +15 °C). Tallene "
    "kan aflæses i SRO eller på databladet. Modellen regner selv COP ud. "
    "Navnet skal stå præcis som i arket Anlaeg. Når arket Varmepumpe er "
    "udfyldt, skal maks varme og COP i Anlaeg stå tomme for varmepumpen — "
    "udfyld kun maks varme, hvis noget andet end varmepumpen selv begrænser "
    "varmen, fx pumper eller nettilslutning. Ellers klipper tallet toppen af "
    "sommerydelsen, uden at resultatet ser forkert ud.",
    BODY, wrap=True)
ws.row_dimensions[21].height = 114

sat(ws, "B23", "Farvekoder", H2)
sat(ws, "C24", "Gul celle = du skal udfylde den.", BODY).fill = UDFYLD
sat(ws, "C25", "Grå række = eksempel. Slet den, eller skriv hen over den.", BODY).fill = EKSEMPEL
sat(ws, "C26", "Blå tekst = tal, du selv taster.", INPUT_F)
sat(ws, "C27", "Enhedstype, balancemarked, DMI-område og priszone har en "
               "dropdown — klik i cellen, og brug pilen til højre. Vises pilen "
               "ikke i din Excel, så skriv værdien af præcis som i listen under "
               "skemaet.", BODY, wrap=True)
ws.row_dimensions[27].height = 42

sat(ws, "B29", "Hvad du IKKE skal udfylde", H2)
sat(ws, "C30",
    "Elspotpriser, balancemarkedspriser og vejrdata henter modellen selv. Du "
    "skal hverken have en Energinet-konto eller en API-nøgle.", BODY, wrap=True)
ws.row_dimensions[30].height = 28

sat(ws, "B32", "Spørgsmål inden dagen", H2)
sat(ws, "C33", "Steen Kramer Jensen, chefkonsulent, Dansk Fjernvarme. "
              "skj@danskfjernvarme.dk", BODY)

# ==============================================================================
# ARK 2 — TIMEDATA
# ==============================================================================
ws = wb.create_sheet("Timedata")
sat(ws, "A1", "Varmeproduktion ab værk — ét års timeværdier", H1)

sat(ws, "A2", "tidszone:", BODY_B)
sat(ws, "B2", "UTC", INPUT_F, UDFYLD, border=True)
sat(ws, "C2", "Skriv enten  UTC  eller  dansk lokaltid  — intet andet. "
              "UTC er det, vi beder om.", BODY_I)

sat(ws, "A3", "årsproduktion ab værk:", BODY_B)
sat(ws, "B3", None, INPUT_F, UDFYLD, border=True).number_format = "0.0"
sat(ws, "C3", "GWh/år, samlet ab værk for 1. juli 2025 – 30. juni 2026. "
              "UDFYLD ALTID — også når timedataene er der. Uden timeserie er "
              "det eneste, modellen har.", BODY_I)

sat(ws, "A4",
    "Indsæt dine data fra række 6. Tidsstempel i kolonne A, MW i kolonne B. "
    "Lad manglende timer stå tomme — skriv ikke nul. Vi regner på "
    "1. juli 2025 – 30. juni 2026.", BODY_I)

header_raekke(ws, 5, ["tidsstempel", "varme_mw", "(noter)"], [26, 14, 62])

start = datetime(2025, 7, 1, 0, 0)
for i, v in enumerate([18.4, 19.1, 19.8, 20.2, 20.0, 19.3, 17.6, 16.9]):
    r = 6 + i
    c = ws.cell(row=r, column=1, value=(start + timedelta(hours=i)))
    c.number_format = "YYYY-MM-DD HH:MM"
    c.font, c.fill, c.border = INPUT_F, EKSEMPEL, BOX
    c = ws.cell(row=r, column=2, value=v)
    c.number_format = "0.00"
    c.font, c.fill, c.border = INPUT_F, EKSEMPEL, BOX
ws.cell(row=6, column=3,
        value="EKSEMPELRÆKKER — slet dem, eller skriv hen over").font = BODY_I

for r in range(14, 44):
    for col in (1, 2):
        c = ws.cell(row=r, column=col)
        c.fill, c.border, c.font = UDFYLD, BOX, INPUT_F
    ws.cell(row=r, column=1).number_format = "YYYY-MM-DD HH:MM"
    ws.cell(row=r, column=2).number_format = "0.00"

ws.freeze_panes = "A6"

# ==============================================================================
# ARK 3 — ANLÆG
# ==============================================================================
ws = wb.create_sheet("Anlaeg")
sat(ws, "A1", "Produktionsenheder", H1)
sat(ws, "A2",
    "Én række pr. enhed. De syv grå rækker viser hver sin type — slet dem, du "
    "ikke har, og skriv dine egne tal i dem, du har. Enhedsnavnet er frit. Du "
    "må gerne have flere enheder af samme slags, og du må gerne forenkle og "
    "lægge fx tre gaskedler sammen til én række.", BODY_I)

kol = ["navn", "type", "maks varme (MW)", "min varme (MW)",
       "varme virkningsgrad", "COP ved 0 °C", "elvirkningsgrad",
       "D&V (kr/MWh)", "startomkostning (kr)", "min driftstid (t)",
       "min stoptid (t)", "kan byde i balancemarked",
       "solvarme: årsproduktion (GWh)"]
header_raekke(ws, 4, kol, [18, 17, 13, 13, 12, 12, 13, 14, 13, 12, 12, 16, 17])

eks = [
    ["varmepumpe",     "heat_pump",      None, 0.0, None, None, None,  30.0,    0,  1, 1, "ja",  None],
    ["elkedel",        "electric_boiler", 4.0, 0.0, 0.99, None, None, 10.0,    0,  1, 1, "ja",  None],
    ["halmkedel",      "biomass_boiler", 12.0, 4.0, 0.96, None, None, 30.0, 8000,  4, 6, "nej", None],
    ["gaskedel",       "gas_boiler",     20.0, 0.0, 0.95, None, None, 15.0,  500,  1, 1, "nej", None],
    ["gasmotor",       "gas_engine_chp", 10.5, 0.9, 0.49, None, 0.41, 100.0, 2000,  2, 2, "ja",  None],
    ["solvarme",       "solar_thermal",  22.0, 0.0, None, None, None,   0.0,    0,  1, 1, "nej", 12.0],
    ["overskudsvarme", "waste_heat",      4.0, 4.0, 1.0,  None, None,   5.0,    0,  1, 1, "nej", None],
]
for i, raekke in enumerate(eks):
    r = 5 + i
    for j, v in enumerate(raekke, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.fill, c.border = INPUT_F, EKSEMPEL, BOX

for r in range(12, 24):
    for j in range(1, 14):
        c = ws.cell(row=r, column=j)
        c.fill, c.border, c.font = UDFYLD, BOX, INPUT_F

dropdown(ws, ENHEDSTYPER, "B5:B23", "Enhedstype")
dropdown(ws, ["ja", "nej"], "L5:L23", "Balancemarked")

sat(ws, "A26", "Lovlige værdier i kolonne 'type' (også i dropdown)", H2)
for i, t in enumerate(ENHEDSTYPER):
    sat(ws, f"A{27 + i}", t, MONO, NOTE_FILL, border=True)
forklar_type = [
    "Varmepumpe. Udfyld arket 'Varmepumpe' med varmeproduktion og eloptag ved "
    "to-tre udetemperaturer, og lad maks varme og COP stå tomme her. "
    "Elvirkningsgrad skal ikke udfyldes. Har I en anden kilde end luft — "
    "grundvand, spildevand, sø — så skriv det i mailen.",
    "Elkedel. Udfyld varme virkningsgrad (MWh varme pr. MWh el, typisk 0,98–0,99).",
    "Halm- eller fliskedel. Brændslet vælges efter, hvilken pris du har udfyldt.",
    "Gaskedel.",
    "Gasmotor med kraftvarme. Udfyld BÅDE varme virkningsgrad og "
    "elvirkningsgrad — begge pr. MWh brændsel. Modellen regner selv "
    "el-til-varme-forholdet ud.",
    "Solvarme. Udfyld kolonnen længst til højre; resten af rækken kan stå som her.",
    "Overskudsvarme fra industri e.l. Sæt min varme = maks varme, hvis leverancen er fast.",
]
for i, tekst in enumerate(forklar_type):
    sat(ws, f"C{27 + i}", tekst, BODY, wrap=True)

sat(ws, "A35", "Kolonnen 'kan byde i balancemarked': skriv  ja  eller  nej", H2)
sat(ws, "C35", "Kun 'ja', hvis enheden faktisk er prækvalificeret hos Energinet. "
               "Er I i tvivl, så skriv nej. I dag byder modellen kun med "
               "varmepumper og elkedler. Et 'ja' ved en gasmotor bliver gemt, "
               "men bruges ikke endnu.", BODY, wrap=True)
ws.row_dimensions[35].height = 42

sat(ws, "A37", "Hvad de øvrige felter betyder", H2)
felter = [
    ("maks varme (MW)", "Enhedens største varmeeffekt ifølge typeskiltet — "
                        "den effekt, den kan levere til nettet ved fuld last. "
                        "Modellen bruger det som et loft, ikke som et mål. "
                        "Varmepumper: lad feltet stå tomt, når arket "
                        "'Varmepumpe' er udfyldt. Udfyld det kun, hvis noget "
                        "andet end varmepumpen selv begrænser varmen — ellers "
                        "klipper det sommerydelsen. Se boksen nedenfor om solvarme."),
    ("min varme (MW)", "Laveste stabile last, når enheden kører. 0, hvis den "
                       "kan regulere helt ned. Bruges kun, når enheden har en "
                       "min driftstid over 1 time eller en min varme over 0."),
    ("varme virkningsgrad", "MWh varme pr. MWh brændsel. Fx 0,95 — ikke 95. "
                            "Kræves for kedler, gasmotor og overskudsvarme. "
                            "For elkedler er 'brændslet' el, typisk 0,98–0,99."),
    ("COP ved 0 °C", "Kun varmepumper, og kun hvis I ikke kan udfylde arket "
                     "'Varmepumpe'. Årsvirkningsgraden skal IKKE bruges her. "
                     "Med kun dette tal regner modellen med fast varmeeffekt "
                     "hele året, og det er mindre præcist."),
    ("elvirkningsgrad", "Kun gasmotor: MWh el pr. MWh brændsel, fx 0,41. "
                        "Sammen med varme virkningsgraden giver den "
                        "el-til-varme-forholdet, modellen regner med. Lad "
                        "feltet stå tomt for alle andre typer."),
    ("min driftstid/stoptid", "Timer enheden mindst skal køre eller stå, når den "
                              "først har skiftet. Kedler med lang opstart: 4–72 timer."),
    ("solvarme: årsproduktion", "Kun for solar_thermal. Hvor mange GWh solfangerne "
                                "leverer på et år. Maks varme er nameplate; profilen "
                                "binder reelt langt lavere."),
]
for i, (navn, tekst) in enumerate(felter, start=38):
    sat(ws, f"A{i}", navn, BODY_B)
    sat(ws, f"C{i}", tekst, BODY, wrap=True)
    ws.row_dimensions[i].height = 30

ws.merge_cells("A46:M46")
boks = sat(ws, "A46",
    "Solvarme og maks varme. For solvarme skriver du nameplate — den effekt, "
    "solfangerfeltet er dimensioneret til. Men solen leverer sjældent i "
    "nærheden af den: indfaldsvinkel, skydække og temperaturtab trækker ned. "
    "Modellen bruger derfor den laveste af to værdier hver time: maks varme og "
    "den genererede solprofil. I eksemplet er nameplate 22 MW, mens profilen "
    "topper omkring 8 MW. Ser du 22 MW i skemaet og højst 8 MW i resultatet, er "
    "det ikke en fejl — det er profilen, der binder. Det er årsproduktionen i "
    "kolonnen længst til højre, der bestemmer, hvor meget sol modellen får.",
    BODY, NOTE_FILL, wrap=True)
ws.row_dimensions[46].height = 88

sat(ws, "A48", "Akkumuleringstanke", H1)
sat(ws, "A49",
    "Én række pr. tank. Har I flere tanke, der arbejder sammen som ét lager, "
    "må I gerne lægge dem sammen til én række.", BODY_I)
header_raekke(ws, 51, ["tank", "volumen (m³)", "maks fyldning (MWh)",
                       "maks ladeeffekt (MW)", "maks afladeeffekt (MW)"],
              [18, 16, 20, 18, 18])
for i, raekke in enumerate([["tank_stor", 7000, 244.0, 25.0, 25.0],
                            ["tank_lille", 3000, 105.0, 25.0, 25.0]], start=52):
    for j, v in enumerate(raekke, start=1):
        c = ws.cell(row=i, column=j, value=v)
        c.font, c.fill, c.border = INPUT_F, EKSEMPEL, BOX
for r in range(54, 58):
    for j in range(1, 6):
        c = ws.cell(row=r, column=j)
        c.fill, c.border, c.font = UDFYLD, BOX, INPUT_F

ws.merge_cells("A59:M60")
sat(ws, "A59",
    "Maks fyldning er den varmemængde, tanken kan rumme fra bund til top — "
    "det tal, driften kender som tankens kapacitet. Kender I den ikke, kan I "
    "regne den: volumen × temperaturspring × 1,163 ÷ 1000. En tank på "
    "7.000 m³ med 30 K mellem frem og retur rummer 244 MWh. Modellen bruger "
    "MWh-tallet; volumen står med, fordi det bruges til at tjekke, at de to "
    "tal passer sammen. Vi spørger ikke om temperaturer.",
    BODY, NOTE_FILL, wrap=True)
ws.row_dimensions[59].height = 46

# ==============================================================================
# ARK 4 — VARMEPUMPE
# ==============================================================================
ws = wb.create_sheet("Varmepumpe")
sat(ws, "A1", "Varmepumper — ydelse ved forskellige udetemperaturer", H1)
ws.merge_cells("A2:E2")
sat(ws, "A2",
    "En luft/vand-varmepumpe leverer mere varme, når det er lunt, og mindre, "
    "når det er koldt. Skriv derfor, hvad varmepumpen leverer i varme og "
    "trækker i el ved fuld last ved mindst to udetemperaturer — gerne tre: en "
    "kold dag (omkring −10 °C), omkring 0 °C og en varm dag (omkring +15 °C). "
    "Tallene kan aflæses i SRO eller på databladet. Modellen regner selv COP "
    "ud og regner mellem punkterne.", BODY_I, wrap=True)
ws.row_dimensions[2].height = 72
ws.merge_cells("A3:E3")
sat(ws, "A3",
    "Én række pr. målepunkt. Navnet skal stå præcis som i arket Anlaeg. Har I "
    "flere varmepumper, så giv hver sine rækker. Grå rækker er eksempel — "
    "slet dem, eller skriv hen over.", BODY_I, wrap=True)
ws.row_dimensions[3].height = 36

header_raekke(ws, 5, ["navn", "udetemperatur (°C)", "varmeproduktion (MW)",
                      "eloptag (MW)", "COP (regnes selv)"],
              [18, 16, 20, 14, 16])
vp_eks = [["varmepumpe", -10, 3.0, 1.25],
          ["varmepumpe", 0, 4.0, 1.38],
          ["varmepumpe", 16, 5.2, 1.53]]
for r in range(6, 24):
    eksempel = r < 6 + len(vp_eks)
    for j in range(1, 5):
        c = ws.cell(row=r, column=j,
                    value=vp_eks[r - 6][j - 1] if eksempel else None)
        c.font, c.border = INPUT_F, BOX
        c.fill = EKSEMPEL if eksempel else UDFYLD
        if j in (3, 4):
            c.number_format = "0.00"
    c = ws.cell(row=r, column=5,
                value=f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),D{r}>0),C{r}/D{r},"")')
    c.font, c.border = BODY, BOX
    c.number_format = "0.00"

ws.merge_cells("A25:E26")
sat(ws, "A25",
    "Tjek COP i kolonne E. En luft/vand-varmepumpe ligger typisk på 2–4, "
    "lavest i kulde. Står der et tal langt fra det, er varme og el byttet om, "
    "eller et af tallene står i kW. Eksemplet viser et typisk forløb: varmen "
    "stiger fra 3,0 til 5,2 MW fra frost til sommer, mens eloptaget kun stiger "
    "lidt.", BODY, NOTE_FILL, wrap=True)
ws.row_dimensions[25].height = 30
ws.row_dimensions[26].height = 30
ws.freeze_panes = "A6"

# ==============================================================================
# ARK 5 — PRISER
# ==============================================================================
ws = wb.create_sheet("Priser")
ws.sheet_view.showGridLines = False
for col, br in zip("ABCDE", (30, 16, 12, 48, 3)):
    ws.column_dimensions[col].width = br

sat(ws, "A1", "Priser, afgifter og tarif", H1)
sat(ws, "A2", "Alt ekskl. moms. Lad felter stå tomme for brændsler, du ikke "
              "bruger. CO2 opgives pr. ton CO2 — ikke pr. MWh gas.", BODY_I)

sat(ws, "A4", "Brændselspriser", H2)
header_raekke(ws, 5, ["brændsel", "pris", "enhed", "note"], [30, 16, 12, 48])
for i, (navn, pris, enhed, note) in enumerate([
    ("naturgas", 606.0, "kr/MWh", "Pr. MWh brændsel, ikke pr. MWh varme. "
     "Indeholder jeres gaspris allerede CO2-afgiften, så skriv 0 i CO2-rækken."),
    ("halm", 232.0, "kr/MWh", "Din indkøbspris ab værk."),
    ("flis", 232.0, "kr/MWh", "Har du både halm og flis, så udfyld begge."),
    ("overskudsvarme", 50.0, "kr/MWh", "Din afregningspris til leverandøren."),
    ("CO2-kvote eller CO2-afgift", 670.0, "kr/t CO2",
     "90 EUR/t × 7,45 kr/EUR. Modellen ganger selv med 0,2 t CO2 pr. MWh gas."),
], start=6):
    sat(ws, f"A{i}", navn, BODY, border=True)
    sat(ws, f"B{i}", pris, INPUT_F, UDFYLD, border=True).number_format = "#,##0.00"
    sat(ws, f"C{i}", enhed, BODY, border=True)
    sat(ws, f"D{i}", note, BODY_I, border=True)

sat(ws, "A12", "Elafgift og faste tarifled", H2)
header_raekke(ws, 13, ["post", "kr/MWh", "", "note"], [30, 16, 12, 48])
for i, (navn, v, note) in enumerate([
    ("elafgift efter elvarmegodtgørelse", 4.0,
     "Det tal, I faktisk betaler — ikke den fulde sats."),
    ("Energinet (transmission+system+balance)", 115.0, "Står samlet på jeres elregning."),
    ("netselskabets drift og vedligehold", 0.0, "Ofte nul. Skriv 0, ikke tomt."),
    ("indfødningstarif (hvis I leverer el)", 4.1,
     "Kun relevant med gasmotor eller anden elproduktion."),
], start=14):
    sat(ws, f"A{i}", navn, BODY, border=True)
    sat(ws, f"B{i}", v, INPUT_F, UDFYLD, border=True).number_format = "#,##0.00"
    sat(ws, f"C{i}", "", BODY, border=True)
    sat(ws, f"D{i}", note, BODY_I, border=True)

sat(ws, "A19", "Tidsvarierende nettarif", H2)
sat(ws, "A20",
    "De tre bånd står på dit netselskabs prisblad, typisk som lavlast, højlast "
    "og spidslast. Har dit selskab andre navne eller flere bånd, så skriv dem, "
    "du har — resten kan stå tomme. Er satserne forskellige vinter og sommer, "
    "så udfyld begge kolonner. Tidspunkterne er faste i modellen og "
    "følger inddelingen i noterne til højre. Bruger jeres netselskab andre "
    "tidspunkter — fx spidslast kl. 17–21 — så skriv jeres tidspunkter i "
    "mailen, når du sender arket.", BODY_I, wrap=True)
ws.row_dimensions[20].height = 72
header_raekke(ws, 21, ["bånd", "vinter (kr/MWh)", "sommer (kr/MWh)", "note"],
              [30, 16, 16, 48])
for i, (navn, vinter, sommer, note) in enumerate([
    ("lavlast", 7.4, 7.4, "Kl. 00–06 alle dage. Om sommeren (april–september) "
                          "også hele weekenden."),
    ("højlast", 14.8, 14.8, "Vinter (oktober–marts): kl. 21–24 på hverdage og "
                            "kl. 06–24 i weekenden. Sommer: kl. 06–24 på hverdage."),
    ("spidslast", 29.5, None, "Kun vinterhverdage kl. 06–21. Har I ikke "
                              "spidslast om sommeren, så lad cellen stå tom."),
], start=22):
    sat(ws, f"A{i}", navn, BODY, border=True)
    for kolonne, v in (("B", vinter), ("C", sommer)):
        sat(ws, f"{kolonne}{i}", v, INPUT_F, UDFYLD, border=True).number_format = "#,##0.00"
    sat(ws, f"D{i}", note, BODY_I, wrap=True, border=True)
for r in (22, 23, 24):
    ws.row_dimensions[r].height = 30

sat(ws, "A26", "Vejrstation, priszone og værkets navn", H2)
header_raekke(ws, 27, ["felt", "værdi", "", "note"], [30, 16, 12, 48])
for i, (felt, vaerdi, note) in enumerate([
    ("nærmeste DMI-område", "karup",
     "Skriv præcis én af:  fyn   vestkyst   karup   ringsted. Vælg det "
     "område, der ligger nærmest værket. Øst for Storebælt: ringsted."),
    ("priszone", "DK1", "Skriv  DK1  vest for Storebælt, ellers  DK2"),
    ("værkets navn", "Andeby Fjernvarme", "Bruges som filnavn for din modelfil."),
], start=28):
    sat(ws, f"A{i}", felt, BODY, border=True)
    sat(ws, f"B{i}", vaerdi, INPUT_F, UDFYLD, border=True)
    sat(ws, f"C{i}", "", BODY, border=True)
    sat(ws, f"D{i}", note, BODY_I, border=True)

ws["D28"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[28].height = 30

DMI_OMRAADER = ["fyn", "vestkyst", "karup", "ringsted"]
dropdown(ws, DMI_OMRAADER, "B28", "DMI-område")
dropdown(ws, ["DK1", "DK2"], "B29", "Priszone")

# Skriver som standard skabelonen dér, hvor vaerksark_til_yaml.py's docstring
# siger den ligger — doc/ i repoets rod, uanset hvor scriptet kaldes fra.
# En sti kan gives som første argument.
import sys
from pathlib import Path

if len(sys.argv) > 1:
    ud = Path(sys.argv[1])
else:
    ud = Path(__file__).resolve().parent.parent / "doc" / "vaerksdata_skabelon.xlsx"
ud.parent.mkdir(parents=True, exist_ok=True)
wb.save(ud)
print(f"skrevet: {ud}")
