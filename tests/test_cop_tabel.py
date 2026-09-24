"""Varmepumpens ydelse som tabel over udetemperaturen (september 2026).

Baggrund: John (Billund) fandt, at modellen havde varmepumpen sat forkert
op. Med type 'linear' ligger varmeloftet fast på p_max_heat, og elforbruget
regnes som varme/COP — så eloptaget er størst i kulde. En rigtig luft/vand-
varmepumpe gør det omvendte: varmen falder i kulde, og eloptaget ændrer sig
kun lidt. Billunds målte punkter:

    -10 °C  12 MW varme  5,0 MW el   COP 2,40
      0 °C  16 MW varme  5,5 MW el   COP 2,91
    +16 °C  21 MW varme  6,2 MW el   COP 3,39

Testene holder fast i de fejl, der er stille:

* et fast varmeloft, der ser rigtigt ud, fordi 16 MW er et rimeligt tal
* COP interpoleret direkte i stedet for varme og el hver for sig
* et ark, hvor navnet i Varmepumpe ikke matcher Anlaeg, og punkterne bare
  forsvinder
* et tal i kW, der giver en COP på 0,003 uden at nogen opdager det

Ingen solver, intet netværk. Modeltestene bygger kun modellen og læser
variablens øvre grænse.
"""
from __future__ import annotations

import datetime as dt
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
import xarray as xr
import yaml

from src.config import COPCurve, load_case

REPO_ROOT = Path(__file__).resolve().parent.parent
BILLUND = [
    {"t_ambient": -10, "heat_mw": 12.0, "el_mw": 5.0},
    {"t_ambient": 0, "heat_mw": 16.0, "el_mw": 5.5},
    {"t_ambient": 16, "heat_mw": 21.0, "el_mw": 6.2},
]


# ------------------------------------------------------------------ COPCurve
def test_tabel_cop_er_varme_over_el_i_punkterne():
    k = COPCurve(type="table", points=BILLUND)
    assert k.evaluate(-10) == pytest.approx(2.40)
    assert k.evaluate(0) == pytest.approx(16 / 5.5)
    assert k.evaluate(16) == pytest.approx(21 / 6.2)


def test_tabel_interpolerer_varme_og_el_hver_for_sig():
    """COP mellem punkterne er varme(T)/el(T), ikke en interpoleret COP.

    Ved 8 °C: varme 18,5, el 5,85 → 3,162. En lineært interpoleret COP
    ville give 3,15. Forskellen er lille, men kun den ene er fysik.
    """
    k = COPCurve(type="table", points=BILLUND)
    assert k.heat_capacity(8) == pytest.approx(18.5)
    assert k.el_capacity(8) == pytest.approx(5.85)
    assert k.evaluate(8) == pytest.approx(18.5 / 5.85)


def test_tabel_holder_yderpunkterne():
    k = COPCurve(type="table", points=BILLUND)
    assert k.heat_capacity(-25) == pytest.approx(12.0)
    assert k.heat_capacity(30) == pytest.approx(21.0)
    assert k.max_el_mw == pytest.approx(6.2)


def test_tabel_sorterer_punkterne():
    k = COPCurve(type="table", points=list(reversed(BILLUND)))
    assert [p["t_ambient"] for p in k.points] == [-10, 0, 16]


def test_tabel_bevarer_xarray_koordinater():
    idx = pd.date_range("2026-01-01", periods=3, freq="h")
    t = xr.DataArray([-10.0, 0.0, 16.0], coords=[("time", idx)])
    ud = COPCurve(type="table", points=BILLUND).heat_capacity(t)
    assert isinstance(ud, xr.DataArray)
    assert (ud.time.values == idx.values).all()
    np.testing.assert_allclose(ud.values, [12.0, 16.0, 21.0])


@pytest.mark.parametrize("punkter, fejl", [
    (BILLUND[:1], "mindst to"),
    ([BILLUND[0], dict(BILLUND[0])], "samme udetemperatur"),
    ([BILLUND[0], {"t_ambient": 0, "heat_mw": 16.0, "el_mw": 5500.0}], "kW"),
    ([BILLUND[0], {"t_ambient": 0, "heat_mw": 5.5, "el_mw": 16.0}], "byttet om"),
    ([BILLUND[0], {"t_ambient": 0, "heat_mw": 0.0, "el_mw": 5.5}], "> 0"),
    ([BILLUND[0], {"t_ambient": 0, "heat_mw": 16.0}], "el_mw"),
])
def test_tabel_afviser_ugyldige_punkter(punkter, fejl):
    with pytest.raises(ValueError, match=fejl):
        COPCurve(type="table", points=punkter)


def test_linear_er_uaendret():
    k = COPCurve()
    assert k.evaluate(0) == pytest.approx(2.2)
    assert k.evaluate(-20) == pytest.approx(1.8)
    assert not k.is_table
    with pytest.raises(ValueError):
        k.heat_capacity(0)


# ---------------------------------------------------------------- modellen
def _andeby_kort(kurve: COPCurve | None, p_max_heat: float = 5.0):
    from src.data_loader import generate_dummy_data

    cfg = load_case(str(REPO_ROOT / "cases" / "andeby.yaml"))
    cfg.time.start = dt.datetime(2026, 1, 1, tzinfo=dt.timezone.utc)
    cfg.time.end = dt.datetime(2026, 1, 1, 5, tzinfo=dt.timezone.utc)
    cfg.units["solvarme"].enabled = False      # kræver profil-CSV
    vp = cfg.units["vp_luft_vand"]
    vp.p_max_heat = p_max_heat
    # Andeby har selv en tabel fra september 2026; None = den gamle lineære.
    vp.cop_curve = kurve if kurve is not None else COPCurve()
    data = generate_dummy_data(cfg)
    data["t_ambient"] = ("time", np.array([-20.0, -10.0, -5.0, 0.0, 16.0, 30.0]))
    return cfg, data


ANDEBY_PUNKTER = [
    {"t_ambient": -10, "heat_mw": 3.0, "el_mw": 1.25},
    {"t_ambient": 0, "heat_mw": 4.0, "el_mw": 1.38},
    {"t_ambient": 16, "heat_mw": 5.2, "el_mw": 1.53},
]


def test_varmeloftet_foelger_udetemperaturen():
    from src.model import build_model

    cfg, data = _andeby_kort(COPCurve(type="table", points=ANDEBY_PUNKTER),
                             p_max_heat=5.0)
    m = build_model(cfg, data)
    loft = m.variables["heat_prod"].upper.sel(unit="vp_luft_vand").values
    # -20 holdes på -10-punktet; 30 °C ville give 5,2, men p_max_heat 5,0 binder.
    np.testing.assert_allclose(loft, [3.0, 3.0, 3.5, 4.0, 5.0, 5.0])


def test_linear_giver_stadig_fast_loft():
    from src.model import build_model

    cfg, data = _andeby_kort(None, p_max_heat=3.4)
    m = build_model(cfg, data)
    loft = m.variables["heat_prod"].upper.sel(unit="vp_luft_vand").values
    np.testing.assert_allclose(loft, 3.4)


def test_tabel_afviser_nan_i_temperaturen():
    from src.model import build_model

    cfg, data = _andeby_kort(COPCurve(type="table", points=ANDEBY_PUNKTER))
    data["t_ambient"] = ("time", np.array([0.0, np.nan, 0.0, 0.0, 0.0, 0.0]))
    with pytest.raises(ValueError, match="NaN"):
        build_model(cfg, data)


def test_reservationsloft_er_tabellens_stoerste_eloptag():
    from src.balancing import _p_el_max

    cfg, data = _andeby_kort(COPCurve(type="table", points=ANDEBY_PUNKTER))
    assert _p_el_max(cfg.units["vp_luft_vand"], data) == pytest.approx(1.53)

    cfg, data = _andeby_kort(None, p_max_heat=3.4)
    # Gammel udledning: p_max_heat / min COP over perioden (clip 1,8 i frost).
    assert _p_el_max(cfg.units["vp_luft_vand"], data) == pytest.approx(3.4 / 1.8)


# ----------------------------------------------------------- konverteringen
def _byg_og_fyld(tmp_path: Path, rediger=None) -> Path:
    import openpyxl

    skabelon = tmp_path / "skabelon.xlsx"
    subprocess.run([sys.executable, str(REPO_ROOT / "scripts" / "byg_skabelon.py"),
                    str(skabelon)], check=True, capture_output=True)
    wb = openpyxl.load_workbook(skabelon)
    ws = wb["Timedata"]
    ws["B3"] = 150
    t = pd.date_range("2025-07-01", "2026-06-30 23:00", freq="h")
    v = 17 + 8 * np.cos((t.dayofyear - 15) / 365 * 2 * np.pi)
    for i, (a, b) in enumerate(zip(t, v)):
        ws.cell(row=6 + i, column=1, value=a.to_pydatetime())
        ws.cell(row=6 + i, column=2, value=round(float(b), 2))
    if rediger:
        rediger(wb)
    ark = tmp_path / "vaerk.xlsx"
    wb.save(ark)
    return ark


def _konverter(tmp_path: Path, ark: Path):
    r = subprocess.run(
        [sys.executable, str(REPO_ROOT / "scripts" / "vaerksark_til_yaml.py"),
         str(ark), "--cases-dir", str(tmp_path / "cases"),
         "--data-dir", str(tmp_path / "data")],
        capture_output=True, text=True, cwd=REPO_ROOT)
    cases = list((tmp_path / "cases").glob("*.yaml"))
    return r, (cases[0] if cases else None)


@pytest.fixture(scope="module")
def modul_tmp(tmp_path_factory):
    return tmp_path_factory.mktemp("vp")


def test_skabelonens_eksempel_giver_tabel(modul_tmp):
    r, case = _konverter(modul_tmp, _byg_og_fyld(modul_tmp))
    assert r.returncode == 0, r.stderr
    vp = yaml.safe_load(case.read_text())["units"]["varmepumpe"]
    assert vp["cop_curve"]["type"] == "table"
    assert [p["t_ambient"] for p in vp["cop_curve"]["points"]] == [-10, 0, 16]
    assert vp["p_max_heat"] == pytest.approx(5.2)      # tom i Anlaeg → tabellens maks
    cfg = load_case(str(case))                          # casen skal kunne læses
    assert cfg.units["varmepumpe"].cop_curve.is_table


def test_uden_maalepunkter_falder_tilbage_paa_cop(tmp_path):
    def uden(wb):
        for r in (6, 7, 8):
            for c in "ABCD":
                wb["Varmepumpe"][f"{c}{r}"] = None
        wb["Anlaeg"]["C5"] = 4.0
        wb["Anlaeg"]["F5"] = 2.9

    r, case = _konverter(tmp_path, _byg_og_fyld(tmp_path, uden))
    assert r.returncode == 0, r.stderr
    assert "overvurderer eloptaget i kulde" in r.stdout
    vp = yaml.safe_load(case.read_text())["units"]["varmepumpe"]
    assert vp["cop_curve"]["type"] == "linear"


def test_ark_fra_foer_v4_uden_varmepumpeark(tmp_path):
    def gammelt(wb):
        del wb["Varmepumpe"]
        wb["Anlaeg"]["C5"] = 4.0
        wb["Anlaeg"]["F5"] = 2.9

    r, case = _konverter(tmp_path, _byg_og_fyld(tmp_path, gammelt))
    assert r.returncode == 0, r.stderr
    assert yaml.safe_load(case.read_text())["units"]["varmepumpe"]["cop_curve"]["type"] == "linear"


def test_navn_der_ikke_matcher_stopper(tmp_path):
    def forkert(wb):
        for r in (6, 7, 8):
            wb["Varmepumpe"][f"A{r}"] = "vp2"

    r, case = _konverter(tmp_path, _byg_og_fyld(tmp_path, forkert))
    assert r.returncode == 1
    assert "'vp2'" in r.stderr and case is None


def test_kw_i_stedet_for_mw_stopper(tmp_path):
    def kw(wb):
        wb["Varmepumpe"]["D7"] = 1380

    r, case = _konverter(tmp_path, _byg_og_fyld(tmp_path, kw))
    assert r.returncode == 1
    assert "kW" in r.stderr and case is None


def test_et_enkelt_punkt_stopper(tmp_path):
    def et(wb):
        for r in (7, 8):
            for c in "ABCD":
                wb["Varmepumpe"][f"{c}{r}"] = None

    r, case = _konverter(tmp_path, _byg_og_fyld(tmp_path, et))
    assert r.returncode == 1
    assert "kun ét målepunkt" in r.stderr


def test_maks_varme_under_tabellen_advarer_med_graense(tmp_path):
    """Typeskiltets tal i Anlaeg klipper sommerydelsen. Det skal siges højt."""
    def typeskilt(wb):
        wb["Anlaeg"]["C5"] = 4.0       # tabellen når 5,2 MW ved +16 °C

    r, case = _konverter(tmp_path, _byg_og_fyld(tmp_path, typeskilt))
    assert r.returncode == 0, r.stderr
    assert "ADVARSEL" in r.stdout and "over ca. 0 °C" in r.stdout
    assert yaml.safe_load(case.read_text())["units"]["varmepumpe"]["p_max_heat"] == 4.0
