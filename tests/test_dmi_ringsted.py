"""DMI-området `ringsted` (30/9 2026).

Skabelonen tilbød `ringsted` i sin dropdown, mens `config.KENDTE_DMI_OMRAADER`
og konverterens egen liste kun kendte fyn, vestkyst og karup. En deltager øst
for Storebælt fik derfor to modstridende beskeder: arket sagde "ringsted", og
konverteren svarede "vælg fyn, vestkyst eller karup".

Testene holder de tre lister sammen (config, konverter, skabelonens dropdown),
og den sidste tester, at loaderen faktisk kan læse `dmi/ringsted_*.csv` fra
df-data.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
import yaml

from src.config import DataOptions, KENDTE_DMI_OMRAADER, load_case
from tests.test_cop_tabel import _byg_og_fyld, _konverter

REPO_ROOT = Path(__file__).resolve().parent.parent
DF_DATA = REPO_ROOT / "data" / "df-data"


def _saet_omraade(omraade: str, zone: str = "DK1"):
    """Skriv DMI-område og priszone i arkets Priser-blok."""
    def rediger(wb):
        ws = wb["Priser"]
        for row in ws.iter_rows():
            for c in row:
                if c.value == "nærmeste DMI-område":
                    ws.cell(row=c.row, column=c.column + 1, value=omraade)
                if c.value == "priszone":
                    ws.cell(row=c.row, column=c.column + 1, value=zone)
    return rediger


def test_ringsted_er_et_kendt_omraade():
    assert "ringsted" in KENDTE_DMI_OMRAADER
    DataOptions(dmi_area="ringsted", price_zone="DK2")     # må ikke kaste


def test_konverteren_accepterer_ringsted(tmp_path):
    ark = _byg_og_fyld(tmp_path, _saet_omraade("ringsted", "DK2"))
    r, case = _konverter(tmp_path, ark)
    assert r.returncode == 0, r.stdout + r.stderr
    d = yaml.safe_load(case.read_text())
    assert d["data"]["dmi_area"] == "ringsted"
    assert d["data"]["price_zone"] == "DK2"
    cfg = load_case(str(case))                 # casen skal kunne indlæses
    assert cfg.data.dmi_area == "ringsted"


def test_ukendt_omraade_stopper_og_navngiver_alle_kendte(tmp_path):
    """Beskeden bygges af config-listen, så den aldrig igen kan udelade et område."""
    ark = _byg_og_fyld(tmp_path, _saet_omraade("bornholm"))
    r, case = _konverter(tmp_path, ark)
    assert r.returncode != 0
    assert case is None
    tekst = r.stdout + r.stderr
    assert "bornholm" in tekst
    for omraade in KENDTE_DMI_OMRAADER:
        assert omraade in tekst, f"beskeden nævner ikke {omraade}: {tekst}"


def test_skabelonens_dropdown_er_lig_config(tmp_path):
    """Det skabelonen tilbyder, skal konverteren kunne tage imod, og omvendt."""
    import openpyxl

    ark = _byg_og_fyld(tmp_path)
    ws = openpyxl.load_workbook(ark)["Priser"]
    lister = [dv.formula1.strip('"').split(",")
              for dv in ws.data_validations.dataValidation
              if "B28" in str(dv.sqref)]
    assert len(lister) == 1, "der skal være præcis én dropdown på DMI-feltet"
    assert set(lister[0]) == set(KENDTE_DMI_OMRAADER)


@pytest.mark.skipif(not (DF_DATA / "dmi").exists(),
                    reason="kræver data/df-data (git clone --depth 1 df-data)")
def test_loaderen_laeser_ringsted_fra_df_data():
    """Data og kode passer sammen: en uge midt i deltagernes år, uden huller."""
    from src.data_loader_github import fetch_dmi_obs_github

    idx = pd.date_range("2025-10-20", "2025-10-27", freq="h")   # inkl. tidsskiftet
    t = fetch_dmi_obs_github("temp_mean_past1h", idx, area="ringsted",
                             repo_root=DF_DATA)
    assert len(t) >= len(idx) - 1
    assert t.notna().all()
    assert -20 < t.min() and t.max() < 40
